"""Consolida os documentos fornecidos, sem alterar os originais."""
from pathlib import Path
from copy import deepcopy
import json
import re
import shutil
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
REPO = ROOT.parents[1]
DOWNLOADS = Path.home() / 'Downloads'
BASE = 'GRUPO XX - 2024-2 - João Vitor Duarte Peçanha Rosa'
MODELS = json.loads((ROOT/'fontes/modelos.json').read_text(encoding='utf-8'))
MODEL = MODELS['pgsql']

def source(num=None):
    ending = f' ({num}).docx' if num else '.docx'
    return next(p for p in DOWNLOADS.glob('DOCUMENT*.docx') if p.name == 'DOCUMENTAÇÃO DO PROJETO'+ending)

def setup(doc):
    sec=doc.sections[0]
    sec.page_width=Cm(21); sec.page_height=Cm(29.7)
    sec.top_margin=Cm(2.3); sec.bottom_margin=Cm(2)
    sec.left_margin=Cm(2.5); sec.right_margin=Cm(2)
    normal=doc.styles['Normal']; normal.font.name='Arial'; normal.font.size=Pt(10.5)
    normal.paragraph_format.line_spacing=1.15
    normal.paragraph_format.space_after=Pt(6)
    for name,size in [('Title',26),('Heading 1',16),('Heading 2',13),('Heading 3',11)]:
        st=doc.styles[name]; st.font.name='Arial'; st.font.size=Pt(size); st.font.color.rgb=RGBColor.from_string('183F55')
        st.paragraph_format.keep_with_next=True
        if name=='Heading 1':st.paragraph_format.page_break_before=True
    doc.styles['Caption'].font.size=Pt(9)
    doc.styles['Caption'].font.color.rgb=RGBColor.from_string('4F5F6B')
    header=sec.header.paragraphs[0]
    header.text='AGENDEI  |  DOCUMENTAÇÃO DO PROJETO'
    header.runs[0].font.size=Pt(8)
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Agendei  •  ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    doc.core_properties.title='Agendei — Documentação consolidada do projeto'
    doc.core_properties.author='João Vitor Duarte Peçanha Rosa; Gabriel Lima Maciel; Thayrine de Lira Maciel; Carlos Eduardo Correa Machado; Fabiano Gonçalves'
    doc.core_properties.subject='Requisitos, funcionamento, telas, MER, DER e atividades do grupo'

def table(doc,headers,rows,size=9):
    t=doc.add_table(rows=1, cols=len(headers));t.style='Light Shading Accent 1'
    for c,h in zip(t.rows[0].cells,headers):c.text=h
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for c,v in zip(t.add_row().cells,row):c.text=str(v)
    for row in t.rows:
        cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(3)
                for r in p.runs:r.font.size=Pt(size)
    return t

def para(doc,text,style=None):
    return doc.add_paragraph(text,style=style)

def image(doc,path,maxw=16.5,maxh=20):
    with Image.open(path) as im:w,h=im.size
    width=min(maxw,maxh*w/h)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path),width=Cm(width))
    return p

def landscape(doc):
    sec=doc.add_section(WD_SECTION_START.NEW_PAGE)
    sec.orientation=WD_ORIENT.LANDSCAPE
    sec.page_width=Cm(42);sec.page_height=Cm(29.7)
    sec.top_margin=sec.bottom_margin=Cm(1.5)
    sec.left_margin=sec.right_margin=Cm(1.5)
    return sec

def portrait(doc):
    sec=doc.add_section(WD_SECTION_START.NEW_PAGE)
    sec.orientation=WD_ORIENT.PORTRAIT
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(2.3);sec.bottom_margin=Cm(2)
    sec.left_margin=Cm(2.5);sec.right_margin=Cm(2)
    return sec

CORRECTIONS={
 'Formato válido e unicidade dentro do estabelecimento.':'Formato válido e unicidade global em usuarios.email. A mesma conta pode possuir vínculos com diferentes estabelecimentos.',
 'Credenciais, dados pessoais, perfil, situação e informações utilizadas na autenticação.':'Identidade global, credenciais, dados pessoais e autenticação. O perfil por empresa é definido em vinculos.tipo.',
 'MySQL/MariaDB':'MySQL/MariaDB ou PostgreSQL',
 'Criação da estrutura e dados iniciais.':'Script de criação da estrutura e dados iniciais para MySQL/MariaDB.',
 'O banco de dados agendei utiliza MySQL ou MariaDB e concentra as informações permanentes da aplicação. As tabelas operacionais são relacionadas ao estabelecimento responsável pelos dados, formando a base do isolamento entre empresas.':'O Agendei suporta MySQL/MariaDB e PostgreSQL por meio de PDO e da camada de adaptação SQL. banco.sql contém o esquema MySQL/MariaDB; banco_postgres.sql contém o equivalente PostgreSQL. Ambos possuem 28 tabelas, 290 colunas e 63 chaves estrangeiras declaradas no código versionado. Os dados operacionais são associados ao estabelecimento; a identidade da pessoa permanece global em usuarios.',
 'Depois da verificação inicial, o sistema aplica uma segunda etapa de autenticação. Uma entre três perguntas pessoais é selecionada: nome materno, data de nascimento ou CEP. A sessão somente é concluída quando a resposta correta é informada.':'Após a verificação inicial, a aplicação solicita um segundo fator. Quando o autenticador TOTP está ativo, exige o código do aplicativo. Nas contas que não o ativaram, mantém a verificação acadêmica por nome materno, data de nascimento ou CEP. A sessão somente é concluída após a validação do fator correspondente.',
 'Um usuário pode possuir função de cliente, profissional ou administrador. Cada agendamento relaciona um cliente, um profissional e um serviço. Os profissionais possuem horários de expediente e podem executar vários serviços. O vínculo com o estabelecimento é utilizado nos relacionamentos para impedir que registros de empresas distintas sejam misturados.':'usuarios representa a pessoa com e-mail único na plataforma. vinculos relaciona essa pessoa ao estabelecimento e define o tipo cliente, profissional ou admin. As tabelas clientes, profissionais e administradores complementam cada vínculo, com id_vinculo único em cada tabela de perfil. Uma pessoa pode participar de várias empresas. Os agendamentos ligam cliente, profissional e serviço; FKs compostas reforçam a associação à mesma empresa.',
 'A estrutura inicial do banco é criada por meio do arquivo banco.sql. Em seguida, o processo de instalação pode cadastrar um administrador, profissionais com expediente configurado, um cliente de demonstração e registros necessários para teste. Após a configuração inicial, o instalador deve ser removido do ambiente de produção.':'Em uma base nova, importe banco.sql para MySQL/MariaDB ou banco_postgres.sql para PostgreSQL. Configure o driver e a conexão em config/database.local.php ou pelas variáveis AGENDEI_DB_*. O instalador pode criar dados demonstrativos; deve ser removido do ambiente de produção após o uso. Os scripts de criação contêm comandos de limpeza, portanto não substituem migrações de uma base existente.',
 'O isolamento também é reforçado por relacionamentos de banco de dados que utilizam chaves associadas ao estabelecimento. Cada empresa possui endereço público próprio, identificado por um slug, e pode personalizar elementos visuais como nome, logotipo, cores e fonte. O usuário cadastrado em determinado endereço permanece vinculado àquela empresa.':'O isolamento é reforçado por chaves compostas associadas ao estabelecimento. Cada empresa possui endereço público identificado por slug e pode personalizar nome, logotipo, cores e fonte. O cadastro cria um vínculo local; a identidade global do usuário pode ser reutilizada em outras empresas mediante autenticação.',
 'O conflito é calculado por profissional. Isso significa que dois profissionais diferentes podem atender no mesmo horário sem que um atendimento bloqueie a agenda do outro.':'A disponibilidade é calculada por profissional, permitindo atendimentos simultâneos em equipes diferentes. Além disso, a criação verifica se o mesmo cliente já possui outro atendimento ativo sobreposto, evitando duas reservas conflitantes para esse cliente.',
 '·                    fila de lembretes com mensagem preparada para WhatsApp;':'· fila de lembretes para WhatsApp, com modo manual e envio automático condicionado à configuração do provedor;',
 '·                    envio automático de mensagens pela API oficial do WhatsApp;':'· ampliar e validar em produção a integração de WhatsApp já implementada, conforme a configuração do provedor;',
 '·                    envio de confirmações e lembretes por e-mail;':'· ampliar os eventos de confirmação e lembretes por e-mail, utilizando a infraestrutura de envio já existente;',
 'Com a utilização de PHP, MySQL ou MariaDB, HTML, CSS e JavaScript, o projeto demonstra a aplicação prática de conceitos de desenvolvimento web e banco de dados em uma necessidade comum a diversos tipos de estabelecimento. O resultado é uma plataforma que oferece maior autonomia ao cliente e maior controle administrativo ao prestador de serviços.':'Com PHP, PDO, MySQL/MariaDB ou PostgreSQL, HTML, CSS e JavaScript, o projeto aplica conceitos de desenvolvimento web e banco de dados à organização de atendimentos, oferecendo autonomia ao cliente e controle ao estabelecimento.',
}

def fix(text):
    text=CORRECTIONS.get(text,text)
    if text.startswith('·'):text=re.sub(r'^·\s*','',text)
    text=re.sub(r'^(\d+\.)\s{3,}',r'\1 ',text)
    return text

doc=Document();setup(doc)
doc.add_paragraph('AGENDEI',style='Title')
doc.add_paragraph('Sistema de agendamento de serviços',style='Subtitle')
doc.add_paragraph('Documentação consolidada do projeto',style='Subtitle')
para(doc,'Front-end • Back-end • Requisitos • MER e DER • Registros do grupo')
para(doc,'Grupo XX — número a preencher')
para(doc,'Integrantes:')
names=['João Vitor Duarte Peçanha Rosa','Gabriel Lima Maciel','Thayrine de Lira Maciel','Carlos Eduardo Correa Machado','Fabiano Gonçalves']
for n in names:para(doc,n)
para(doc,'Registros de desenvolvimento: agosto a outubro de 2026')
para(doc,'Revisão documental: 10 de outubro de 2026')
doc.add_page_break()
para(doc,'SUMÁRIO',style='Title')
toc=OxmlElement('w:fldSimple');toc.set(qn('w:instr'),'TOC \\o "1-2" \\h \\z \\u')
doc.add_paragraph()._p.append(toc)

doc.add_heading('1 Apresentação da entrega',1)
para(doc,'Este documento reúne o relatório do sistema, a análise de requisitos, as telas públicas capturadas, o modelo de dados e os registros de reuniões e atividades fornecidos pelo grupo. As informações técnicas foram revisadas com base nos arquivos versionados do Agendei.')
table(doc,['Critério de avaliação','Localização'],[
 ('Requisitos funcionais e não funcionais','Capítulo 2: requisitos numerados e regras de negócio.'),
 ('Identidade visual, telas, validações e responsividade','Capítulos 3 e 4: descrição e capturas reais de telas públicas.'),
 ('Modelo MER e DER','Capítulo 5 e anexos: modelo conceitual, diagramas físicos e dicionário.'),
 ('Reuniões, decisões e atividades por componente','Capítulo 6: cronologia e contribuições individuais registradas.'),
 ('Front-end, PHP e criação do banco','Capítulo 7 e pacote ZIP: código da aplicação e scripts SQL.'),
])
para(doc,'A nomenclatura 2024-2 do arquivo foi mantida conforme o enunciado fornecido. Os registros de reuniões dos documentos de origem são de 2026. O número do grupo não foi informado e permanece indicado por XX.')

doc.add_heading('2 Análise de requisitos',1)
doc.add_heading('2.1 Atores',2)
table(doc,['Ator','Responsabilidade'],[
 ('Visitante','Consultar a apresentação pública da plataforma e acessar os formulários de entrada e cadastro.'),
 ('Cliente','Consultar serviços e disponibilidade, reservar, acompanhar e cancelar os próprios atendimentos.'),
 ('Profissional','Consultar a própria agenda e executar operações autorizadas de atendimento e bloqueio.'),
 ('Administrador do estabelecimento','Gerenciar os dados operacionais e as configurações da empresa.'),
 ('Administrador master global','Administrar estabelecimentos e contas globais em área separada.'),
])
req=Document(source(2))
rf=[];rn=[];rnf=[]
updates={
 'RF01':'Visualizar a página inicial: apresentar o produto Agendei, seus recursos e os acessos de cadastro e login. O link de um estabelecimento resolve sua entrada local; o catálogo para agendamento é consultado na área do cliente.',
 'RF02':'Cadastrar cliente: permitir cadastro com nome, CPF, nascimento, sexo, nome materno, contatos, endereço, login, senha e confirmação.',
 'RF03':'Validar cadastro: conferir os campos no navegador e no servidor; e-mail é único na plataforma, enquanto CPF e login possuem restrições por estabelecimento. Contas existentes utilizam o fluxo de vínculo.',
 'RF04':'Realizar login: na entrada global, aceitar e-mail e senha; na entrada da empresa, aceitar e-mail ou login local e senha. Concluir o acesso após o segundo fator aplicável.',
 'RF25':'Criar bloqueio de agenda: permitir que o profissional bloqueie períodos da própria agenda quando autorizado pela configuração do estabelecimento e por sua permissão individual.',
 'RF26':'Remover bloqueio: permitir a exclusão de bloqueios da própria agenda segundo as permissões verificadas no servidor.',
 'RN07':'O mesmo cliente não pode possuir agendamentos ativos com intervalos sobrepostos, mesmo com profissionais diferentes.',
 'RNF02':'Tecnologia: utilizar PHP 8 ou superior, PDO, HTML, CSS e JavaScript, com suporte a MySQL/MariaDB ou PostgreSQL.',
 'RNF10':'Integridade dos dados: utilizar PKs, FKs, restrições de unicidade, índices e transações. Referências lógicas sem FK devem ser identificadas como tais no modelo.',
 'RNF17':'Desempenho: empregar índices e filtros; aplicar paginação nas listagens que a implementam. Não se estabelece aqui uma meta de tempo de resposta medida.',
}
for p in req.paragraphs:
    match=re.match(r'^(RF\d+|RNF\d+|RN\d+)\s*[—:]\s*(.+)',p.text)
    if match:
        key,text=match.groups();text=updates.get(key,text)
        (rnf if key.startswith('RNF') else rf if key.startswith('RF') else rn).append((key,text))
rf.extend([
 ('RF41','Gerenciar vínculos: permitir que uma pessoa utilize a mesma identidade em diferentes estabelecimentos e selecione o contexto autorizado de acesso.'),
 ('RF42','Autenticar em segunda etapa: validar código TOTP quando ativado ou o desafio cadastral aplicável; limitar tentativas.'),
 ('RF43','Administrar a plataforma: disponibilizar área master separada para gestão de estabelecimentos, contas e auditoria global.'),
 ('RF44','Gerenciar filiais: associar unidades, profissionais e reservas conforme as regras implementadas.'),
 ('RF45','Gerenciar diferenciais: disponibilizar, quando habilitados, espera, pacotes, fidelidade, avaliações e sinal Pix com confirmação manual.'),
 ('RF46','Gerenciar lembretes: manter fila e estado de envio; usar modo manual ou provedor configurado para WhatsApp.'),
])
rn.extend([
 ('RN13','Identidade global e vínculo local: usuarios.email é único na plataforma; vinculos.tipo define o perfil em cada estabelecimento.'),
 ('RN14','Preservar dados da reserva: valor e hora_fim são gravados na criação e não recalculados por alterações posteriores do catálogo.'),
 ('RN15','Cancelar preserva a reserva: alterar o status, registrar motivo/responsável quando aplicáveis e liberar o intervalo.'),
 ('RN16','Crédito de pacote: id_cliente_pacote não possui FK física; a aplicação valida titularidade, serviço, saldo, validade e pagamento.'),
])
doc.add_heading('2.2 Requisitos funcionais',2);table(doc,['Código','Requisito'],rf)
doc.add_heading('2.3 Regras de negócio',2);table(doc,['Código','Regra'],rn)
doc.add_heading('2.4 Requisitos não funcionais',2);table(doc,['Código','Requisito'],rnf)
doc.add_heading('2.5 Rastreabilidade das funcionalidades',2)
table(doc,['Requisitos','Arquivos de referência'],[
 ('RF01–RF07, RF41–RF42','index.php; cadastro.php; entrar.php; login.php; vincular.php; dois_fatores.php; includes/auth.php'),
 ('RF08–RF18','cliente/dashboard.php; cliente/agendar.php; cliente/agendamentos.php; cliente/historico.php; cliente/perfil.php'),
 ('RF19–RF27','profissional/dashboard.php; profissional/agenda.php; profissional/horarios.php; profissional/bloqueios.php; profissional/perfil.php'),
 ('RF28–RF40','admin/dashboard.php; admin/servicos.php; admin/profissionais.php; admin/horarios.php; admin/agendamentos.php; admin/relatorios.php; admin/configuracoes.php'),
 ('RF43–RF46','master/; admin/filiais.php; admin/diferenciais.php; models/Diferencial.php; models/Lembrete.php; models/WhatsApp.php'),
 ('RN01–RN10, RN14–RN16','models/Agendamento.php; models/Disponibilidade.php; models/Diferencial.php'),
 ('RNF03–RNF11','config/database.php; includes/seguranca.php; includes/auth.php; models/Sql.php; banco.sql; banco_postgres.sql'),
])

doc.add_heading('3 Relatório técnico do sistema',1)
report=Document(source(1)); started=False; ended=False
for node in report.element.body:
    if node.tag==qn('w:p'):
        p=Paragraph(node,report);text=p.text.strip()
        if text=='1 INTRODUÇÃO':started=True
        if not started:continue
        if text=='REFERÊNCIAS':break
        if not text:continue
        if p.style.name.startswith('Heading'):
            match=re.match(r'^(\d+(?:\.\d+)*)\s+(.+)$',text)
            if match:
                number,title=match.groups()
                level=2 if '.' not in number else 3
                doc.add_heading('3.'+number+' '+title.capitalize(),level)
        else:
            para(doc,fix(text),'List Bullet' if text.startswith('·') else None)
    elif node.tag==qn('w:tbl') and started:
        src=Table(node,report);rows=[[fix(c.text) for c in r.cells] for r in src.rows]
        if rows[0][0]=='Tabela':
            rows.insert(3,['vinculos','Associação da pessoa ao estabelecimento, com perfil, status, login e último acesso locais.'])
        if rows[0][0]=='Diretório/arquivo':rows.append(['banco_postgres.sql','Esquema equivalente de criação para PostgreSQL.'])
        if rows[0][0]=='Comando':rows.append(['php tests/modelo_bd.php','Verificar cobertura e cardinalidades do modelo sem acesso ao banco.'])
        table(doc,rows[0],rows[1:])

doc.add_heading('4 Interface, telas e validações',1)
para(doc,'As figuras deste capítulo foram capturadas da aplicação PHP executada localmente em 10/10/2026, usando navegador Edge em modo automatizado. Foram utilizadas resoluções de 1366 × 900 e 390 × 844 pixels. As capturas mostram telas públicas reais; não foram criadas representações dos painéis autenticados.')
doc.add_heading('4.1 Identidade visual e responsividade',2)
para(doc,'A página do produto apresenta a marca Agendei, navegação, chamadas de cadastro e entrada e seus recursos. A entrada global utiliza e-mail; a entrada local da empresa também pode aceitar o login do vínculo. O cadastro organiza dados pessoais, contato, endereço e credenciais, com rótulos e orientações de preenchimento.')
figs=[
 ('inicio_desktop_recorte.png','Página inicial do produto em computador.'),
 ('entrada_desktop_recorte.png','Entrada global: formulário de e-mail e senha.'),
 ('cadastro_desktop_recorte.png','Cadastro do cliente: organização dos dados pessoais.'),
 ('cadastro_empresa_desktop_recorte.png','Cadastro público de estabelecimento.'),
]
for i,(file,title) in enumerate(figs,1):
    doc.add_page_break();image(doc,ROOT/'imagens'/file,maxh=19)
    para(doc,f'Figura {i} — {title}',style='Caption')
doc.add_page_break()
pair=doc.add_table(rows=1,cols=2)
for c,file in zip(pair.rows[0].cells,['inicio_celular_recorte.png','entrada_celular_recorte.png']):
    c.paragraphs[0].add_run().add_picture(str(ROOT/'imagens'/file),width=Cm(7.2))
para(doc,'Figura 5 — Página inicial e entrada global em celular (390 × 844).',style='Caption')
para(doc,'Nas oito combinações de tela e resolução verificadas, a largura do documento coincidiu com a largura da viewport. Essa medição verifica ausência de rolagem horizontal nessas telas e resoluções; não constitui uma auditoria de todas as páginas ou dispositivos.')
doc.add_heading('4.2 Críticas dos dados preenchidos',2)
table(doc,['Campo','Exemplo inválido','Resposta prevista'],[
 ('Nome','Ana','Indicar o limite de 15 a 80 caracteres e o uso de letras e espaços.'),
 ('CPF','111.111.111-11','Rejeitar sequência repetida e informar CPF inválido.'),
 ('E-mail','texto sem formato de e-mail','Informar erro de formato; o servidor também verifica a unicidade global.'),
 ('Login','abc','Exigir exatamente seis letras no cadastro local.'),
 ('Senha','abc123','Exigir a regra acadêmica de oito letras no cadastro do cliente.'),
 ('Confirmação','Valor diferente da senha','Informar que as senhas não conferem.'),
 ('Reserva','Intervalo já ocupado','Revalidar no servidor e impedir uma reserva conflitante.'),
])
image(doc,ROOT/'imagens/cadastro_erros.png',maxh=18)
para(doc,'Figura 6 — Mensagens reais de crítica de cadastro com nome curto, CPF repetido e campos obrigatórios incompletos.',style='Caption')
para(doc,'O evento de validação local destacou 17 campos inválidos e impediu o envio do formulário. Não foi criado um cadastro. A implementação PHP repete as verificações no servidor, independentemente do JavaScript.')
doc.add_heading('4.3 Telas autenticadas previstas no projeto',2)
table(doc,['Perfil','Telas implementadas de referência'],[
 ('Cliente','Dashboard, agendamento por etapas, listagem de reservas, histórico e perfil.'),
 ('Profissional','Dashboard, agenda, expediente, bloqueios e perfil.'),
 ('Administrador','Dashboard, agenda, serviços, profissionais, clientes, horários, relatórios, configurações e logs.'),
 ('Master global','Dashboard, estabelecimentos, contas, segurança e auditoria.'),
])
para(doc,'As telas autenticadas estão descritas e identificadas no código, mas suas capturas não foram incluídas nesta revisão. A demonstração acadêmica deve utilizar uma base de teste para apresentar reserva, conflito, cancelamento e permissões por perfil.')

doc.add_heading('5 MER e DER',1)
doc.add_heading('5.1 Modelo conceitual',2)
para(doc,'O núcleo do modelo separa pessoa, empresa e participação. usuarios identifica a pessoa; estabelecimento identifica a empresa; vinculos relaciona ambos e define o perfil local. clientes, profissionais e administradores complementam o vínculo. A conta master global é independente dessa especialização.')
para(doc,'O catálogo reúne serviços e profissionais. profissional_servico resolve a associação N:N entre eles. Expedientes e bloqueios determinam disponibilidade. agendamentos relaciona cliente, profissional e serviço e preserva data, intervalo e valor da reserva. Filiais organizam unidades. Pacotes, pagamentos, fidelidade, espera, avaliações e notificações complementam a operação.')
table(doc,['Relacionamento','Cardinalidade e interpretação'],[
 ('usuarios → vinculos','Uma pessoa pode ter zero ou vários vínculos; cada vínculo possui uma pessoa.'),
 ('estabelecimento → vinculos','Uma empresa pode ter zero ou vários vínculos; cada vínculo pertence a uma empresa.'),
 ('vinculos → tabelas de perfil','Um vínculo tem zero ou um registro em cada tabela de perfil, conforme a unicidade de id_vinculo. A aplicação seleciona o perfil compatível com vinculos.tipo.'),
 ('profissionais ↔ servicos','N:N resolvida por profissional_servico; a associação usa chave primária composta.'),
 ('clientes/profissionais/servicos → agendamentos','Cada reserva exige um cliente, um profissional e um serviço. Cada participante pode possuir zero ou várias reservas.'),
 ('filiais → profissionais/agendamentos','A filial é opcional no registro operacional; uma filial pode possuir zero ou vários registros.'),
 ('estabelecimento → estabelecimento_plano/assinaturas','Cada empresa possui zero ou um registro em cada tabela; id_estabelecimento também é a PK dessas tabelas.'),
])
doc.add_heading('5.2 Modelo físico e legenda',2)
para(doc,'Os DERs revisados foram gerados por código a partir do esquema PostgreSQL versionado, incluindo colunas e FKs adicionadas por ALTER TABLE. Não representam inspeção de uma base instalada. Os dois scripts possuem 28 tabelas, 290 colunas e 63 FKs declaradas; algumas FKs simples reforçadas por compostas aparecem agrupadas apenas no desenho.')
para(doc,'PK = chave primária; FK = participação em chave estrangeira física; UK = participação em restrição de unicidade, que pode ser composta. A marca UK em uma coluna não significa que essa coluna seja única isoladamente. || = exatamente um; círculo e barra = zero ou um; círculo e pé de galinha = zero ou vários. Linhas contínuas representam relações identificadoras; tracejadas representam relações não identificadoras. Ambas correspondem a FKs físicas.')
para(doc,'Os diagramas por módulo mostram todas as PKs e FKs das entidades selecionadas e resumem os atributos descritivos. Entidades de referência podem aparecer em mais de um módulo. Relações que atravessam módulos devem ser consultadas no DER completo e no dicionário de dados. As páginas dos diagramas usam A3 em paisagem para facilitar a leitura.')
doc.add_heading('5.3 Referências lógicas sem FK',2)
table(doc,['Referência','Tratamento'],[
 ('agendamentos.id_cliente_pacote → cliente_pacotes','Não possui FOREIGN KEY nos scripts; saldo, validade, titularidade e serviço são conferidos pela aplicação.'),
 ('logs_autenticacao.id_usuario → usuarios','Identificador histórico sem FK; o log pode sobreviver à exclusão da pessoa.'),
 ('logs_master.id_master → administradores_master','Referência histórica sem FK.'),
 ('logs_master.id_estabelecimento → estabelecimento','Referência histórica sem FK.'),
 ('sessoes_lembradas.id_estabelecimento','Coluna sem FK física. A tabela possui FKs declaradas para id_usuario e id_vinculo.'),
])
para(doc,'bloqueios_agenda.id_usuario_criou e agendamentos.id_usuario_cancelou possuem FKs reais para usuarios e foram corrigidas nos diagramas. id_cliente_pacote permanece sem marcação FK. A tabela sessoes_lembradas não recebeu uma FK de empresa inexistente.')
modules=[('01_nucleo','Núcleo, empresas e perfis'),('02_disponibilidade','Catálogo e disponibilidade'),('03_reservas','Reserva e seus participantes'),('04_atendimento','Espera, notificações e avaliações'),('05_beneficios','Pacotes, pagamentos e fidelidade'),('06_plataforma','Planos e assinatura da empresa'),('07_seguranca','Acesso, sessões e auditoria')]
for i,(name,title) in enumerate(modules,1):
    landscape(doc)
    p=doc.add_heading(f'5.{i+3} DER — {title}',2);p.paragraph_format.page_break_before=False
    image(doc,ROOT/'imagens'/f'der_{name}.png',maxw=39,maxh=23)
    para(doc,f'DER físico PostgreSQL por módulo. Fonte: banco_postgres.sql; tipos resumidos. Dicionário completo em anexo.',style='Caption')
portrait(doc)

doc.add_heading('6 Gestão e atividades do grupo',1)
doc.add_heading('6.1 Integrantes e responsabilidades registradas',2)
table(doc,['Integrante','Área','Atividades descritas nos registros'],[
 (names[0],'Back-end','Rotas e endpoints; operações de agendamento, alteração e cancelamento; testes e revisão de funcionalidades.'),
 (names[1],'Front-end','Navegação, componentes visuais, formulários e integração das telas com a API.'),
 (names[2],'Back-end','Regras de negócio, validações, prevenção de conflitos, tratamento de erros e testes.'),
 (names[3],'Back-end','Estrutura do servidor, gerenciamento dos dados, integração com o banco e revisão final.'),
 (names[4],'Front-end','Criação das telas e do layout, ajustes visuais e finalização da interface.'),
])
para(doc,'As contribuições acima e os relatos seguintes reproduzem as informações fornecidas pelo grupo. Não foram deduzidos autores a partir do código ou do histórico Git.')
doc.add_heading('6.2 Reuniões em ordem cronológica',2)
meeting=Document(source());paras=[p.text for p in meeting.paragraphs if p.text.strip()]
records=[]
for i,text in enumerate(paras):
    match=re.match(r'\d+\. Reunião — (\d{2}/\d{2}/\d{4})',text)
    if match:records.append((match.group(1),'17h',paras[i+1]))
first=next(i for i,t in enumerate(paras) if t.startswith('PRIMEIRA REUNIÃO'))
second=next(i for i,t in enumerate(paras) if t.startswith('SEGUNDA REUNIÃO'))
end=next(i for i,t in enumerate(paras) if t.startswith('RELATÓRIO DE REUNIÕES'))
records.extend([('27/08/2026','Não informado',' '.join(paras[first+1:second])),('03/09/2026','Não informado',' '.join(paras[second+1:end]))])
records.sort(key=lambda x:tuple(reversed(x[0].split('/'))))
para(doc,'Os documentos de origem continham uma lista semanal e dois relatos datados de 27/08 e 03/09, chamados de primeira e segunda reunião. A consolidação ordena todos os registros pela data e remove essa numeração conflitante. Mantém 17h somente nas reuniões em que esse horário foi informado.')
for date,time,text in records:
    doc.add_heading(f'{date} — {time}',3)
    para(doc,text)
para(doc,'O relato de 27/08 informa a formalização da inclusão de Fabiano, enquanto a lista semanal já o cita em atividades anteriores. Ambos foram preservados como registros de origem; esta revisão não determina a data efetiva de ingresso. O intervalo declarado no relatório vai até 07/10, mas a última reunião datada fornecida é 03/10; não foi acrescentada reunião em 07/10.')
doc.add_heading('6.3 Cronograma de referência',2)
chron=Document(source(3))
for t in chron.tables:
    rows=[[c.text for c in r.cells] for r in t.rows]
    table(doc,rows[0],rows[1:])
para(doc,'O cronograma foi preservado como planejamento de oito semanas, sem atribuir datas ou comprovação de execução não presentes no documento original.')

doc.add_heading('7 Código, banco e instalação',1)
table(doc,['Componente','Arquivos incluídos no pacote'],[
 ('Front-end','Páginas PHP que produzem HTML; assets/css/; assets/js/; assets/img/; includes/ de layout.'),
 ('Back-end','PHP da raiz; admin/; api/; cliente/; profissional/; master/; models/; includes/; config/.'),
 ('Criação do banco','banco.sql (MySQL/MariaDB) e banco_postgres.sql (PostgreSQL).'),
 ('Atualização de instalações existentes','scripts/migrar*.php e migrations/.'),
 ('Modelo','DER completo .mmd e .svg; sete diagramas por módulo; dicionário PostgreSQL.'),
 ('Documentação','Este documento em .docx, .doc e PDF; LEIAME.md e guias do repositório.'),
])
para(doc,'Para uma instalação nova: preparar PHP e a extensão PDO do banco escolhido; importar somente o SQL correspondente; configurar AGENDEI_DB_DRIVER e as demais variáveis de conexão ou o arquivo local; iniciar o servidor web; executar a configuração demonstrativa apenas em ambiente de teste. O esquema PostgreSQL deve ser importado em um banco previamente criado, pois não contém CREATE DATABASE nem USE.')
para(doc,'Os arquivos de criação contêm comandos DROP TABLE e não devem ser importados em uma base com dados a preservar. Atualizações de uma instalação existente devem utilizar as migrações previstas no projeto. O pacote omite credenciais locais, arquivos .env, logs de execução e dados de usuários.')

doc.add_heading('8 Verificação realizada nesta revisão',1)
table(doc,['Verificação','Resultado','Limite'],[
 ('php tests/modelo_bd.php','598 verificações passaram.','Cobertura do esquema e cardinalidades; não conecta ao banco.'),
 ('php tests/requisitos.php','65 verificações; zero falhas.','Regras de validação e autenticação verificadas sem banco.'),
 ('Páginas públicas','Oito capturas em computador/celular; HTTP 200.','Página inicial, entrada global e dois cadastros.'),
 ('Largura responsiva','Sem transbordamento horizontal nas oito combinações.','1366 e 390 pixels; não abrange todas as telas.'),
 ('Crítica de cadastro','Envio impedido; 17 campos destacados.','Validação JavaScript local; nenhuma gravação de cadastro.'),
])
para(doc,'Não foram executados nesta revisão os testes de integração com banco nem a navegação autenticada. As verificações acima comprovam somente os resultados indicados, sem afirmar aprovação de todos os requisitos ou funcionamento dos provedores externos.')

doc.add_heading('Referências e fontes do projeto',1)
for p in report.paragraphs:
    if p.text.startswith(('ASSOCIAÇÃO','BRASIL.','PHP DOCUMENTATION','ORACLE.')):para(doc,p.text)
para(doc,'POSTGRESQL GLOBAL DEVELOPMENT GROUP. Documentação oficial do PostgreSQL. Referência técnica complementar ao esquema banco_postgres.sql.')
para(doc,'FONTES FORNECIDAS PELO GRUPO. DOCUMENTAÇÃO DO PROJETO.docx e partes (1), (2), (3) e (4), além dos três diagramas PNG encaminhados para revisão.')
para(doc,'FONTES LOCAIS. banco.sql; banco_postgres.sql; models/ModeloBanco.php; config/database.php; models/Usuario.php; models/Agendamento.php; includes/auth.php; páginas PHP e assets do repositório Agendei.')

doc.save(ROOT/(BASE+'.docx'))

# Anexo completo de atributos, unicidades e todas as FKs, independente do resumo visual.
dd=Document();setup(dd)
dd.add_paragraph('AGENDEI — Dicionário PostgreSQL',style='Title')
para(dd,'Anexo técnico da documentação consolidada — 10/10/2026')
para(dd,'Fonte: banco_postgres.sql versionado. Contém 28 tabelas, 290 colunas e 63 FKs físicas. A marca UK indica participação em uma chave possivelmente composta; consulte os conjuntos listados. Tipos e definições exatos são preservados abaixo.')
for name,t in MODEL['tabelas'].items():
    dd.add_heading(name,1)
    table(dd,['Coluna','Tipo SQL','Nulo','Chaves'],[(c['nome'],c['tipo'],'Sim' if c['nulo'] else 'Não',', '.join(c['chaves']) or '—') for c in t['colunas'].values()],size=8)
    para(dd,'PK: '+', '.join(t['pk']))
    for k in t['uk']:para(dd,'UK composta/conjunto: ('+', '.join(k)+')')
    if t['fks']:
        table(dd,['FK / colunas locais','Destino','Ações'],[(fk['nome']+'\n('+', '.join(fk['campos'])+')',fk['pai']+' ('+', '.join(fk['referencias'])+')',fk['acoes'] or 'Padrão do banco') for fk in t['fks']],size=8)
    else:para(dd,'Nenhuma FK física declarada nesta tabela.')
    para(dd,'Definições exatas de coluna (incluem valores padrão):')
    for c in t['colunas'].values():
        p=para(dd,c['definicao'])
        p.paragraph_format.space_after=Pt(2)
        for r in p.runs:r.font.name='Consolas';r.font.size=Pt(7.5)
    checks=[r for r in t['restricoes'] if 'CHECK' in r]
    for c in checks:para(dd,c)
dd.add_heading('Cardinalidades de todas as FKs',1)
table(dd,['Restrição','Pai → filha','Pais por filha','Filhas por pai'],[(f['nome'],f['pai']+' → '+f['filha'],f['por_filha'],f['por_pai']) for f in MODEL['relacoes']],size=7.5)
dd.save(ROOT/'ANEXO - Dicionário de dados PostgreSQL.docx')

# Markdown pesquisável e revisável do texto consolidado.
md=[]
for el in doc.element.body:
    if el.tag==qn('w:p'):
        p=Paragraph(el,doc)
        if not p.text.strip():continue
        level=int(p.style.name[-1]) if p.style.name.startswith('Heading') else 0
        md.append(('#'*level+' ' if level else '')+p.text)
    elif el.tag==qn('w:tbl'):
        t=Table(el,doc);rows=[[c.text.replace('\n',' / ').replace('|','/') for c in r.cells] for r in t.rows]
        if rows:
            md.append('| '+' | '.join(rows[0])+' |')
            md.append('| '+' | '.join(['---']*len(rows[0]))+' |')
            md.extend('| '+' | '.join(r)+' |' for r in rows[1:])
    md.append('')
(ROOT/'DOCUMENTACAO_REVISADA.md').write_text('\n'.join(md).rstrip()+'\n',encoding='utf-8')
print('DOCUMENTOS GERADOS:',BASE,flush=True)
